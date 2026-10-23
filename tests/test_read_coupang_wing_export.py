"""Synthetic, in-memory contract tests; no business workbook or fixture is read."""

from __future__ import annotations

import builtins
import hashlib
import io
import json
from datetime import UTC, datetime, timedelta, timezone
from xml.etree.ElementTree import Element, fromstring, tostring
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

import topspeed_ai_os.tools as tools_package
import topspeed_ai_os.tools.read as read_package
import topspeed_ai_os.tools.read.coupang_wing_export as wing_module
from topspeed_ai_os.core.ids import InvalidIdentifierError, generate_run_id, generate_task_id
from topspeed_ai_os.tools import ErrorCategory, ExecutionMode, SensitiveBytes, ToolStatus
from topspeed_ai_os.tools.read import TabularData
from topspeed_ai_os.tools.read.coupang_wing_export import CoupangWingExportReadTool

SALES = "SALES_PERIOD_AGGREGATE"
INVENTORY = "INVENTORY_HEALTH_OBSERVATION"
KST = timezone(timedelta(hours=9))
CLOSE = datetime(2026, 10, 20, tzinfo=KST)
FETCHED = CLOSE + timedelta(days=1)
REQUEST_ID = "req_" + "d" * 32
CREATED_TAG = "{http://purl.org/dc/terms/}created"
SALES_HEADERS = ("등록상품ID", "옵션 ID", "판매방식", "총 판매수", "총 취소된 상품수", "판매량")
INVENTORY_HEADERS = (
    "등록상품 ID",
    "옵션 ID",
    "SKU ID",
    "상품등급",
    "판매가능재고 (실시간 기준)",
    "입고예정재고(실시간 기준)",
    "고객반품 지난 30일(전일자 기준)",
)
SALES_ROW = ("SYNTHETIC-PRODUCT-001", "SYNTHETIC-OPTION-001", "SYNTHETIC-MODE", "17", "-2", "15")
INVENTORY_ROW = (
    "SYNTHETIC-PRODUCT-001",
    "SYNTHETIC-OPTION-001",
    "SYNTHETIC-SKU-001",
    "SYNTHETIC-GRADE",
    "12",
    "4",
    "3",
)
METADATA_KEYS = {
    "schema_version",
    "export_kind",
    "worksheet_name",
    "workbook_created_at",
    "raw_artifact_checksum_sha256",
    "original_source_uri",
    "period_start",
    "period_end",
    "period_timezone",
    "period_close",
    "selection_authority_reference",
    "envelope_sha256",
}
AUTHORITY_REFERENCE = (
    "docs/architecture/"
    "ADR_PHASE5_COUPANG_PERIOD_AGGREGATE_READ_ONLY_INVENTORY_ASSESSMENT_V1.md"
    "@f5952d3dcb808c4cb65e6f18bc64b52afcbeab6d"
)
SALES_AUTHORITY_REFERENCE = (
    "docs/architecture/"
    "ADR_PHASE5_COUPANG_L_REPEATABLE_COMPLETED_30D_READ_ONLY_AUTHORITY_V1.md"
    "@b02fb74a766f4f85e7c7d878a3a5faa4a16bc7ba"
)


def _rewrite_parts(raw, replacements=None, *, remove=(), comment=b""):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(raw)) as source, ZipFile(output, "w", ZIP_DEFLATED) as destination:
        for info in source.infolist():
            if info.filename not in remove:
                destination.writestr(
                    info, (replacements or {}).get(info.filename, source.read(info.filename))
                )
        destination.comment = comment
    return output.getvalue()


def _book(
    kind=SALES,
    *,
    created="2026-10-20T00:00:00+09:00",
    headers=None,
    data_rows=None,
    sheet_name=None,
    explanatory=None,
    first_other_sheet=False,
):
    book = Workbook()
    book.properties.creator = "SYNTHETIC-AUTHOR"
    book.properties.created = datetime(2000, 1, 1)
    book.properties.modified = datetime(2000, 1, 1)
    sheet = book.active
    if first_other_sheet:
        sheet.title = "SYNTHETIC-OTHER"
        sheet.append(["SYNTHETIC-NON-AUTHORITATIVE"])
        sheet = book.create_sheet()
    sheet.title = sheet_name or ("vendor item metrics" if kind == SALES else "Excel Data")
    sheet.append(
        headers if headers is not None else SALES_HEADERS if kind == SALES else INVENTORY_HEADERS
    )
    if kind == INVENTORY:
        sheet.append(
            explanatory
            if explanatory is not None
            else ["SYNTHETIC-EXPLANATION"] * len(INVENTORY_HEADERS)
        )
    for row in (
        data_rows if data_rows is not None else [SALES_ROW if kind == SALES else INVENTORY_ROW]
    ):
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    raw = output.getvalue()
    with ZipFile(io.BytesIO(raw)) as archive:
        core = fromstring(archive.read("docProps/core.xml"))
    for node in list(core):
        if node.tag == CREATED_TAG:
            core.remove(node)
    for value in created if isinstance(created, tuple) else (created,):
        if value is not None:
            node = Element(CREATED_TAG)
            node.text = value
            core.append(node)
    return _rewrite_parts(raw, {"docProps/core.xml": tostring(core)})


def _tool(*, clock=None, now=FETCHED, request_id_factory=None):
    return CoupangWingExportReadTool(
        clock=(lambda: now) if clock is None else clock,
        request_id_factory=(lambda: REQUEST_ID)
        if request_id_factory is None
        else request_id_factory,
    )


def _read(raw, kind=SALES, *, tool=None, **overrides):
    arguments = {"raw_bytes": raw, "export_kind": kind}
    if kind == SALES:
        arguments.update(
            period_start="2026-09-20", period_end="2026-10-19", period_timezone="Asia/Seoul"
        )
    arguments.update(overrides)
    return (_tool() if tool is None else tool).read(**arguments)


def _independent_hash(payload):
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode()
    return "sha256:" + hashlib.sha256(encoded).hexdigest()


def _expected_envelope(raw, kind, created_text):
    sales = kind == SALES
    return {
        "schema_version": (
            "topspeed.coupang_wing_export.source_envelope.v2"
            if sales
            else "topspeed.coupang_wing_export.source_envelope.v1"
        ),
        "source": "coupang_wing_export",
        "source_uri": (
            "coupang-wing://seller-insights/vendor-item-metrics"
            if sales
            else "coupang-wing://inventory-health/sku-info"
        ),
        "export_kind": kind,
        "worksheet_name": "vendor item metrics" if sales else "Excel Data",
        "raw_checksum_sha256": "sha256:" + hashlib.sha256(raw).hexdigest(),
        "workbook_created_at": created_text,
        "period_start": "2026-09-20" if sales else None,
        "period_end": "2026-10-19" if sales else None,
        "period_timezone": "Asia/Seoul" if sales else None,
        "period_close": CLOSE.isoformat() if sales else None,
        "data_as_of": CLOSE.isoformat() if sales else created_text,
        "selection_authority_reference": SALES_AUTHORITY_REFERENCE
        if sales
        else AUTHORITY_REFERENCE,
    }


def _assert_failed(result, category, *, operation):
    assert result.status is ToolStatus.FAILED
    assert result.status is not ToolStatus.PARTIAL
    assert result.tool_name == "coupang_wing_export_read"
    assert result.operation == operation
    assert result.execution_mode is ExecutionMode.MOCK
    assert result.source_metadata is None
    assert result.raw_result is None
    assert result.parsed_result is None
    assert result.error.category is category
    assert result.error.retryable is False
    assert result.error.http_status is None
    assert result.attempt_count == 1


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize(
    ("created", "canonical"),
    [
        ("2026-10-20T00:00:00+09:00", "2026-10-20T00:00:00+09:00"),
        ("2026-10-19T15:00:00Z", "2026-10-20T00:00:00+09:00"),
        ("2026-10-20T01:00:00+10:00", "2026-10-20T00:00:00+09:00"),
        ("2026-10-19T10:00:00-05:00", "2026-10-20T00:00:00+09:00"),
        ("2026-10-20T02:04:06+09:00", "2026-10-20T02:04:06+09:00"),
    ],
)
def test_successful_envelope_has_exact_source_time_bytes_and_independent_hash(
    kind, created, canonical
):
    raw = _book(kind, created=created, first_other_sheet=True)
    result = _read(raw, kind)
    expected = _expected_envelope(raw, kind, canonical)
    metadata = result.source_metadata
    assert result.status is ToolStatus.SUCCESS
    assert result.error is None
    assert result.tool_name == "coupang_wing_export_read"
    assert result.operation == (
        "coupang_wing_export.read_sales_period_aggregate"
        if kind == SALES
        else "coupang_wing_export.read_inventory_health_observation"
    )
    assert result.execution_mode is ExecutionMode.MOCK
    assert result.request_id == REQUEST_ID
    assert result.attempt_count == 1
    assert result.warnings == ()
    assert metadata.source == expected["source"]
    assert metadata.source_uri == expected["source_uri"]
    assert metadata.checksum_sha256 == expected["raw_checksum_sha256"]
    assert metadata.data_as_of == datetime.fromisoformat(expected["data_as_of"])
    assert metadata.fetched_at == FETCHED
    assert (
        metadata.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert metadata.http_status is None
    assert metadata.modified_time is None
    specific = metadata.source_specific_metadata
    assert type(specific) is dict
    assert set(specific) == METADATA_KEYS
    assert specific == {
        "schema_version": expected["schema_version"],
        "export_kind": kind,
        "worksheet_name": expected["worksheet_name"],
        "workbook_created_at": canonical,
        "raw_artifact_checksum_sha256": expected["raw_checksum_sha256"],
        "original_source_uri": expected["source_uri"],
        "period_start": expected["period_start"],
        "period_end": expected["period_end"],
        "period_timezone": expected["period_timezone"],
        "period_close": expected["period_close"],
        "selection_authority_reference": expected["selection_authority_reference"],
        "envelope_sha256": _independent_hash(expected),
    }
    assert type(result.raw_result) is SensitiveBytes
    assert result.raw_result.reveal() == raw
    assert type(result.parsed_result) is TabularData
    assert result.parsed_result.sheet_name == expected["worksheet_name"]
    assert result.parsed_result.headers == (SALES_HEADERS if kind == SALES else INVENTORY_HEADERS)
    assert result.parsed_result.rows == ((SALES_ROW if kind == SALES else INVENTORY_ROW),)
    assert "SYNTHETIC-PRODUCT-001" not in repr(result)
    assert "SYNTHETIC-AUTHOR" not in repr(result)


def test_inventory_explanatory_row_is_not_data():
    raw = _book(INVENTORY, explanatory=["=1+1", "SYNTHETIC-NOT-A-RECORD"])
    result = _read(raw, INVENTORY)
    assert result.status is ToolStatus.SUCCESS
    assert result.parsed_result.rows == (INVENTORY_ROW,)


def test_extra_non_authoritative_column_is_preserved():
    raw = _book(
        headers=(*SALES_HEADERS, "SYNTHETIC-NOTE"), data_rows=[(*SALES_ROW, "SYNTHETIC-TEXT")]
    )
    result = _read(raw)
    assert result.status is ToolStatus.SUCCESS
    assert result.parsed_result.headers[-1] == "SYNTHETIC-NOTE"
    assert result.parsed_result.rows[0][-1] == "SYNTHETIC-TEXT"


def test_byte_difference_changes_raw_and_envelope_hash_without_changing_rows():
    raw = _book()
    changed = _rewrite_parts(raw, comment=b"SYNTHETIC-BYTE-CHANGE")
    first, second = _read(raw), _read(changed)
    assert first.parsed_result == second.parsed_result
    assert first.raw_result.reveal() == raw
    assert second.raw_result.reveal() == changed
    assert first.source_metadata.checksum_sha256 != second.source_metadata.checksum_sha256
    for payload, result in ((raw, first), (changed, second)):
        assert result.source_metadata.source_specific_metadata[
            "envelope_sha256"
        ] == _independent_hash(_expected_envelope(payload, SALES, CLOSE.isoformat()))
    assert (
        first.source_metadata.source_specific_metadata["envelope_sha256"]
        != (second.source_metadata.source_specific_metadata["envelope_sha256"])
    )


@pytest.mark.parametrize("kwargs", [{}, {"clock": None}, {"clock": 1}, {"clock": False}])
def test_constructor_requires_explicit_callable_clock(kwargs):
    with pytest.raises(TypeError):
        CoupangWingExportReadTool(**kwargs)


def test_explicit_falsey_callable_does_not_trigger_wall_clock_fallback():
    class Clock:
        def __bool__(self):
            return False

        def __call__(self):
            return CLOSE

    result = _read(_book(), tool=_tool(clock=Clock()))
    assert result.status is ToolStatus.SUCCESS
    assert result.source_metadata.fetched_at == CLOSE


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize("delta", [timedelta(0), timedelta(hours=1)])
def test_clock_called_once_and_equal_or_later_acquisition_is_valid(kind, delta):
    calls = []

    def clock():
        calls.append(True)
        return CLOSE + delta

    result = _read(_book(kind), kind, tool=_tool(clock=clock))
    assert result.status is ToolStatus.SUCCESS
    assert calls == [True]
    assert result.source_metadata.fetched_at == CLOSE + delta
    assert result.source_metadata.data_as_of == CLOSE


def test_naive_clock_cannot_produce_success():
    with pytest.raises(ValueError, match="timezone-aware"):
        _read(_book(), tool=_tool(now=datetime(2026, 10, 21)))


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
def test_acquisition_before_creation_cannot_produce_success(kind):
    result = _read(_book(kind), kind, tool=_tool(now=CLOSE - timedelta(seconds=1)))
    _assert_failed(
        result,
        ErrorCategory.MALFORMED_SOURCE,
        operation=(
            "coupang_wing_export.read_sales_period_aggregate"
            if kind == SALES
            else "coupang_wing_export.read_inventory_health_observation"
        ),
    )


def test_acquisition_and_request_identity_do_not_change_envelope_hash():
    raw = _book()
    first = _read(raw, tool=_tool(now=CLOSE))
    second = _read(
        raw, tool=_tool(now=FETCHED.astimezone(UTC), request_id_factory=lambda: "req_" + "e" * 32)
    )
    assert first.source_metadata.fetched_at != second.source_metadata.fetched_at
    assert first.request_id != second.request_id
    assert (
        first.source_metadata.source_specific_metadata
        == second.source_metadata.source_specific_metadata
    )


@pytest.mark.parametrize("infrastructure", ["clock", "request_id_factory"])
def test_injected_infrastructure_exceptions_propagate(infrastructure):
    def broken():
        raise RuntimeError("SYNTHETIC-INFRASTRUCTURE-FAILURE")

    with pytest.raises(RuntimeError, match="SYNTHETIC-INFRASTRUCTURE-FAILURE"):
        _read(_book(), tool=_tool(**{infrastructure: broken}))


def test_invalid_request_id_factory_result_preserves_result_validation():
    with pytest.raises(ValueError, match="request_id"):
        _read(_book(), tool=_tool(request_id_factory=lambda: "SYNTHETIC-INVALID-ID"))


@pytest.mark.parametrize("kind", [None, 1, "", "SYNTHETIC-INVALID-KIND", "sales_period_aggregate"])
def test_invalid_export_kind_uses_safe_constant_operation(kind):
    result = _read(_book(), kind)
    _assert_failed(
        result, ErrorCategory.MALFORMED_INPUT, operation="coupang_wing_export.read_invalid_request"
    )
    assert "SYNTHETIC-INVALID-KIND" not in repr(result)


@pytest.mark.parametrize("field", ["period_start", "period_end", "period_timezone"])
@pytest.mark.parametrize("value", [None, 1, "SYNTHETIC-WRONG-SELECTION"])
def test_sales_selection_must_be_exact(field, value):
    result = _read(_book(), **{field: value})
    _assert_failed(
        result,
        ErrorCategory.MALFORMED_INPUT,
        operation="coupang_wing_export.read_sales_period_aggregate",
    )


@pytest.mark.parametrize("field", ["period_start", "period_end", "period_timezone"])
def test_inventory_rejects_any_non_none_period_argument(field):
    result = _read(_book(INVENTORY), INVENTORY, **{field: "SYNTHETIC-NOT-NONE"})
    _assert_failed(
        result,
        ErrorCategory.MALFORMED_INPUT,
        operation="coupang_wing_export.read_inventory_health_observation",
    )


@pytest.mark.parametrize(
    "raw", [None, "SYNTHETIC-NOT-BYTES", bytearray(b"synthetic"), memoryview(b"x"), 1]
)
def test_only_exact_builtin_bytes_are_accepted(raw):
    _assert_failed(
        _read(raw),
        ErrorCategory.MALFORMED_INPUT,
        operation="coupang_wing_export.read_sales_period_aggregate",
    )


def test_subclasses_cannot_spoof_exact_input_types():
    class BytesSubclass(bytes):
        pass

    class StringSubclass(str):
        pass

    raw = _book()
    assert _read(BytesSubclass(raw)).error.category is ErrorCategory.MALFORMED_INPUT
    assert _read(raw, StringSubclass(SALES)).operation == "coupang_wing_export.read_invalid_request"
    assert _read(raw, period_timezone=StringSubclass("Asia/Seoul")).status is ToolStatus.FAILED


@pytest.mark.parametrize("raw", [b"", b"SYNTHETIC-NOT-XLSX", b"PK\x03\x04SYNTHETIC-BROKEN-ZIP"])
def test_malformed_workbooks_fail_without_payload(raw):
    _assert_failed(
        _read(raw),
        ErrorCategory.MALFORMED_SOURCE,
        operation="coupang_wing_export.read_sales_period_aggregate",
    )


@pytest.mark.parametrize(
    "created",
    [
        None,
        ("2026-10-20T00:00:00+09:00", "2026-10-20T00:00:00+09:00"),
        "SYNTHETIC-NOT-A-TIME",
        "2026-10-20T00:00:00",
        "2026-10-20T00:00:00.000+09:00",
        "2026-10-20T00:00:00+09:00:01",
        "2026-10-20T00:00:00+09:00:00.5",
        "2026-10-20T00:00:00+09:60",
        "2026-10-20T00:00:00+24:00",
        "2026-10-20T24:00:00+09:00",
        "2026-02-30T00:00:00+09:00",
        " 2026-10-20T00:00:00+09:00",
    ],
)
def test_creation_metadata_cardinality_and_w3cdtf_are_required(created):
    _assert_failed(
        _read(_book(created=created)),
        ErrorCategory.MALFORMED_SOURCE,
        operation="coupang_wing_export.read_sales_period_aggregate",
    )


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
def test_pre_close_creation_is_ineligible_for_both_exports(kind):
    result = _read(_book(kind, created="2026-10-19T23:59:59+09:00"), kind)
    assert result.status is ToolStatus.FAILED
    assert result.error.category is ErrorCategory.MALFORMED_SOURCE
    assert result.source_metadata is result.raw_result is result.parsed_result is None


def test_missing_and_malformed_core_properties_are_rejected():
    raw = _book()
    for changed in (
        _rewrite_parts(raw, remove=("docProps/core.xml",)),
        _rewrite_parts(raw, {"docProps/core.xml": b"<SYNTHETIC-BROKEN-XML"}),
        _rewrite_parts(raw, {"[Content_Types].xml": b"<SYNTHETIC-BROKEN-XML"}),
    ):
        assert _read(changed).error.category is ErrorCategory.MALFORMED_SOURCE


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize("fault", ["shared_string_index", "style_descriptor"])
def test_malformed_xlsx_decoder_index_and_type_errors_are_source_failures(kind, fault):
    raw = _book(kind)
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    part = "xl/worksheets/sheet1.xml" if fault == "shared_string_index" else "xl/styles.xml"
    with ZipFile(io.BytesIO(raw)) as archive:
        root = fromstring(archive.read(part))
    if fault == "shared_string_index":
        coordinate = "A2" if kind == SALES else "A3"
        cell = next(node for node in root.iter(namespace + "c") if node.get("r") == coordinate)
        cell.clear()
        cell.set("r", coordinate)
        cell.set("t", "s")
        value = Element(namespace + "v")
        value.text = "999999"
        cell.append(value)
    else:
        # A fill without a valid child cannot satisfy openpyxl's typed descriptor.
        next(root.iter(namespace + "fill")).clear()
    changed = _rewrite_parts(raw, {part: tostring(root)})
    _assert_failed(
        _read(changed, kind),
        ErrorCategory.MALFORMED_SOURCE,
        operation=(
            "coupang_wing_export.read_sales_period_aggregate"
            if kind == SALES
            else "coupang_wing_export.read_inventory_health_observation"
        ),
    )


def test_unsupported_workbook_content_type_is_not_accepted_as_xlsx():
    raw = _book()
    with ZipFile(io.BytesIO(raw)) as archive:
        content_types = archive.read("[Content_Types].xml")
    changed = _rewrite_parts(
        raw,
        {
            "[Content_Types].xml": content_types.replace(
                b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            )
        },
    )
    assert _read(changed).error.category is ErrorCategory.UNSUPPORTED_FORMAT


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
def test_missing_required_sheet_has_no_first_sheet_fallback(kind):
    result = _read(_book(kind, sheet_name="SYNTHETIC-WRONG-SHEET"), kind)
    assert result.error.category is ErrorCategory.SOURCE_NOT_FOUND


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize("change", ["missing", "duplicate", "whitespace"])
def test_required_header_change_is_schema_failure(kind, change):
    headers = list(SALES_HEADERS if kind == SALES else INVENTORY_HEADERS)
    row = list(SALES_ROW if kind == SALES else INVENTORY_ROW)
    if change == "missing":
        headers.pop()
        row.pop()
    elif change == "duplicate":
        headers.append(headers[0])
        row.append("SYNTHETIC-DUPLICATE")
    else:
        headers[0] += " "
    result = _read(_book(kind, headers=headers, data_rows=[row]), kind)
    assert result.error.category is ErrorCategory.SOURCE_SCHEMA_CHANGED


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize("position", ["header", "data"])
def test_formulas_in_consumed_cells_are_failures_not_warnings(kind, position):
    headers = list(SALES_HEADERS if kind == SALES else INVENTORY_HEADERS)
    row = list(SALES_ROW if kind == SALES else INVENTORY_ROW)
    (headers if position == "header" else row)[0] = "=1+1"
    result = _read(_book(kind, headers=headers, data_rows=[row]), kind)
    assert result.error.category is ErrorCategory.MALFORMED_SOURCE
    assert result.warnings == ()
    assert result.raw_result is result.parsed_result is result.source_metadata is None


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
@pytest.mark.parametrize("shape", ["short", "wide", "empty"])
def test_row_width_and_empty_body_are_not_hidden_by_dimensions(kind, shape):
    row = SALES_ROW if kind == SALES else INVENTORY_ROW
    rows = [] if shape == "empty" else [row[:-1] if shape == "short" else (*row, "SYNTHETIC-EXTRA")]
    result = _read(_book(kind, data_rows=rows), kind)
    assert result.status is ToolStatus.FAILED
    assert result.error.category is (
        ErrorCategory.MALFORMED_SOURCE if shape == "empty" else ErrorCategory.SOURCE_SCHEMA_CHANGED
    )


def test_valid_correlation_ids_are_preserved_but_not_hashed():
    raw = _book()
    task_id, run_id = generate_task_id(), generate_run_id()
    result = _read(raw, task_id=task_id, run_id=run_id)
    assert result.task_id == task_id
    assert result.run_id == run_id
    assert (
        result.source_metadata.source_specific_metadata
        == _read(raw).source_metadata.source_specific_metadata
    )


@pytest.mark.parametrize("field", ["task_id", "run_id"])
def test_invalid_correlation_ids_keep_existing_exception_behavior(field):
    with pytest.raises(InvalidIdentifierError):
        _read(_book(), **{field: "SYNTHETIC-INVALID-ID"})


def test_unexpected_internal_error_is_redacted(monkeypatch):
    raw = _book()

    def broken_loader(*args, **kwargs):
        raise RuntimeError("SYNTHETIC-PRIVATE-PATH-AUTHOR-AND-PAYLOAD")

    monkeypatch.setattr(wing_module, "load_workbook", broken_loader)
    result = _read(raw)
    _assert_failed(
        result, ErrorCategory.UNKNOWN, operation="coupang_wing_export.read_sales_period_aggregate"
    )
    assert "SYNTHETIC-PRIVATE" not in repr(result)
    assert result.error.redacted_summary == "Source envelope validation failed"


@pytest.mark.parametrize("error_type", [IndexError, TypeError])
def test_injected_internal_index_and_type_errors_remain_unknown(monkeypatch, error_type):
    raw = _book()

    def broken_loader(*args, **kwargs):
        raise error_type("SYNTHETIC-INTERNAL-IMPLEMENTATION-FAULT")

    monkeypatch.setattr(wing_module, "load_workbook", broken_loader)
    result = _read(raw)
    _assert_failed(
        result, ErrorCategory.UNKNOWN, operation="coupang_wing_export.read_sales_period_aggregate"
    )
    assert "SYNTHETIC-INTERNAL" not in repr(result)


@pytest.mark.parametrize("infrastructure", ["clock", "request_id_factory"])
@pytest.mark.parametrize("error_type", [IndexError, TypeError])
def test_infrastructure_index_and_type_errors_still_propagate(infrastructure, error_type):
    def broken():
        raise error_type("SYNTHETIC-INFRASTRUCTURE-FAULT")

    with pytest.raises(error_type, match="SYNTHETIC-INFRASTRUCTURE-FAULT"):
        _read(_book(), tool=_tool(**{infrastructure: broken}))


@pytest.mark.parametrize("kind", [SALES, INVENTORY])
def test_metadata_has_fresh_mutation_isolated_dictionary_per_invocation(kind):
    raw = _book(kind)
    tool = _tool()
    first, second = _read(raw, kind, tool=tool), _read(raw, kind, tool=tool)
    first_dict = first.source_metadata.source_specific_metadata
    second_dict = second.source_metadata.source_specific_metadata
    assert type(first_dict) is type(second_dict) is dict
    assert set(first_dict) == set(second_dict) == METADATA_KEYS
    assert first_dict is not second_dict
    snapshot = dict(second_dict)
    first_dict.clear()
    first_dict["SYNTHETIC-MUTATION"] = "SYNTHETIC-VALUE"
    assert second_dict == snapshot
    assert _read(raw, kind, tool=tool).source_metadata.source_specific_metadata == snapshot


def test_no_package_export_activation():
    for package in (tools_package, read_package):
        assert "CoupangWingExportReadTool" not in package.__all__
        assert not hasattr(package, "CoupangWingExportReadTool")


@pytest.mark.parametrize(
    "field", ["source", "source_uri", "worksheet_name", "data_as_of", "checksum_sha256"]
)
def test_caller_cannot_inject_derived_authority(field):
    with pytest.raises(TypeError):
        _read(_book(), **{field: "SYNTHETIC-INJECTED"})


def test_source_boundary_uses_no_filesystem_network_or_environment(monkeypatch):
    import os
    import socket

    raw = _book()

    def forbidden(*args, **kwargs):
        raise AssertionError("SYNTHETIC-FORBIDDEN-EXTERNAL-ACCESS")

    monkeypatch.setattr(builtins, "open", forbidden)
    monkeypatch.setattr(io, "open", forbidden)
    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr(os, "getenv", forbidden)
    assert _read(raw).status is ToolStatus.SUCCESS


@pytest.mark.parametrize(
    "start,end,close_text",
    [
        ("2026-09-20", "2026-10-19", "2026-10-20T00:00:00+09:00"),
        ("2026-10-02", "2026-10-31", "2026-11-01T00:00:00+09:00"),
        ("2026-12-15", "2027-01-13", "2027-01-14T00:00:00+09:00"),
        ("2028-02-01", "2028-03-01", "2028-03-02T00:00:00+09:00"),
        ("2027-02-01", "2027-03-02", "2027-03-03T00:00:00+09:00"),
        ("2024-02-01", "2024-03-01", "2024-03-02T00:00:00+09:00"),
    ],
)
def test_explicit_completed_period_derives_v2_close_without_date_inference(start, end, close_text):
    close = datetime.fromisoformat(close_text)
    raw = _book(created=close_text)
    calls = []

    def clock():
        calls.append(True)
        return close + timedelta(days=1)

    result = _read(raw, period_start=start, period_end=end, tool=_tool(clock=clock))
    assert result.status is ToolStatus.SUCCESS
    assert calls == [True]
    metadata = result.source_metadata
    expected = {
        **_expected_envelope(raw, SALES, close_text),
        "period_start": start,
        "period_end": end,
        "period_close": close_text,
        "data_as_of": close_text,
    }
    assert metadata.data_as_of == close
    assert set(metadata.source_specific_metadata) == METADATA_KEYS
    assert metadata.source_specific_metadata["schema_version"] == expected["schema_version"]
    assert metadata.source_specific_metadata["selection_authority_reference"] == (
        SALES_AUTHORITY_REFERENCE
    )
    assert metadata.source_specific_metadata["envelope_sha256"] == _independent_hash(expected)


@pytest.mark.parametrize(
    "start,end,zone",
    [
        ("2026-10-02", "2026-10-30", "Asia/Seoul"),  # 29 days
        ("2026-10-01", "2026-10-31", "Asia/Seoul"),  # 31 days
        ("2027-02-29", "2027-03-30", "Asia/Seoul"),
        ("2026-11-31", "2026-12-30", "Asia/Seoul"),
        ("20261002", "20261031", "Asia/Seoul"),
        ("2026-10-2", "2026-10-31", "Asia/Seoul"),
        ("2026-10-02 ", "2026-10-31", "Asia/Seoul"),
        ("2026-10-31", "2026-10-02", "Asia/Seoul"),
        ("2026-10-02", "2026-10-31", "+09:00"),
        ("9999-12-02", "9999-12-31", "Asia/Seoul"),
    ],
)
def test_noncanonical_or_not_exactly_thirty_day_period_fails_closed(start, end, zone):
    result = _read(_book(), period_start=start, period_end=end, period_timezone=zone)
    assert result.status is ToolStatus.FAILED
    assert result.error.category is ErrorCategory.MALFORMED_INPUT
    assert result.raw_result is result.source_metadata is result.parsed_result is None


@pytest.mark.parametrize("offset", [-1, 0])
def test_dynamic_sales_close_boundary(offset):
    close = datetime(2027, 1, 14, tzinfo=KST)
    created = close + timedelta(seconds=offset)
    result = _read(
        _book(created=created.isoformat()),
        period_start="2026-12-15",
        period_end="2027-01-13",
        tool=_tool(now=close),
    )
    assert result.status is (ToolStatus.FAILED if offset < 0 else ToolStatus.SUCCESS)
    if offset < 0:
        assert result.error.category is ErrorCategory.MALFORMED_SOURCE


def test_inventory_retains_v1_identity_and_no_sales_period_binding():
    created = "2027-01-14T00:00:00+09:00"
    result = _read(
        _book(INVENTORY, created=created),
        INVENTORY,
        tool=_tool(now=datetime.fromisoformat(created)),
    )
    specific = result.source_metadata.source_specific_metadata
    assert result.status is ToolStatus.SUCCESS
    assert specific["schema_version"] == "topspeed.coupang_wing_export.source_envelope.v1"
    assert specific["selection_authority_reference"] == AUTHORITY_REFERENCE
    for name in ("period_start", "period_end", "period_timezone", "period_close"):
        assert specific[name] is None
    assert result.source_metadata.data_as_of.isoformat() == created
