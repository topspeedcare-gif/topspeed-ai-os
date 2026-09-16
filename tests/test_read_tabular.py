from __future__ import annotations

import io
from datetime import UTC, datetime

import pytest
from openpyxl import Workbook

from topspeed_ai_os.core.ids import generate_run_id, generate_task_id
from topspeed_ai_os.tools import ErrorCategory, ExecutionMode, ToolStatus, raw_sha256
from topspeed_ai_os.tools.read import FormulaCell, TabularReadTool

NOW = datetime(2026, 9, 16, 0, 0, tzinfo=UTC)
REQUEST_ID = "req_" + ("d" * 32)


@pytest.fixture
def tool():
    return TabularReadTool(
        execution_mode=ExecutionMode.MOCK,
        clock=lambda: NOW,
        request_id_factory=lambda: REQUEST_ID,
    )


def make_xlsx(*, formula=False, extra_sheet=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["sku", "value"])
    sheet.append(["TS-001", "=1+1" if formula else 2])
    if extra_sheet:
        other = workbook.create_sheet("Other")
        other.append(["name", "value"])
        other.append(["x", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("raw", "encoding", "expected"),
    [
        (b"name,value\na,1\n", "utf-8", (("a", "1"),)),
        (b"name,value\r\na,1\r\n", "utf-8", (("a", "1"),)),
        (b"\xef\xbb\xbfname,value\na,1\n", "utf-8-sig", (("a", "1"),)),
        ("name,value\n한글,1\n".encode("cp949"), "cp949", (("한글", "1"),)),
    ],
)
def test_csv_encodings_and_line_endings(tool, raw, encoding, expected):
    result = tool.read(
        raw_bytes=raw,
        source_uri="file:///input.csv",
        file_format="csv",
        encoding=encoding,
    )
    assert result.status is ToolStatus.SUCCESS
    assert result.parsed_result.headers == ("name", "value")
    assert result.parsed_result.rows == expected
    assert result.source_metadata.data_as_of is None
    assert result.source_metadata.checksum_sha256 == raw_sha256(raw)


def test_csv_explicit_delimiter(tool):
    result = tool.read(
        raw_bytes=b"name;value\na;1\n",
        source_uri="file:///input.csv",
        file_format="csv",
        delimiter=";",
    )
    assert result.parsed_result.rows == (("a", "1"),)


@pytest.mark.parametrize(
    ("raw", "encoding", "category"),
    [
        (b"", "utf-8", ErrorCategory.MALFORMED_SOURCE),
        (b"\n", "utf-8", ErrorCategory.MALFORMED_SOURCE),
        (b"name,name\na,b\n", "utf-8", ErrorCategory.MALFORMED_SOURCE),
        (b"name,value\na\n", "utf-8", ErrorCategory.MALFORMED_SOURCE),
        (b"name,value\n\xff,1\n", "utf-8", ErrorCategory.MALFORMED_SOURCE),
    ],
)
def test_csv_malformed_inputs(tool, raw, encoding, category):
    result = tool.read(
        raw_bytes=raw,
        source_uri="file:///input.csv",
        file_format="csv",
        encoding=encoding,
    )
    assert result.status is ToolStatus.FAILED
    assert result.error.category is category


def test_csv_invalid_delimiter_is_malformed_input(tool):
    result = tool.read(
        raw_bytes=b"name,value\na,1\n",
        source_uri="file:///input.csv",
        file_format="csv",
        delimiter="||",
    )
    assert result.error.category is ErrorCategory.MALFORMED_INPUT


def test_xlsx_valid_workbook_and_explicit_sheet(tool):
    raw = make_xlsx(extra_sheet=True)
    result = tool.read(
        raw_bytes=raw,
        source_uri="file:///input.xlsx",
        file_format="xlsx",
        sheet_name="Other",
    )
    assert result.status is ToolStatus.SUCCESS
    assert result.parsed_result.sheet_name == "Other"
    assert result.parsed_result.headers == ("name", "value")
    assert result.parsed_result.rows == (("x", 3),)
    assert result.source_metadata.checksum_sha256 == raw_sha256(raw)


def test_xlsx_formula_is_preserved_and_not_executed(tool):
    raw = make_xlsx(formula=True)
    result = tool.read(
        raw_bytes=raw,
        source_uri="file:///formula.xlsx",
        file_format="xlsx",
    )
    formula = result.parsed_result.rows[0][1]
    assert isinstance(formula, FormulaCell)
    assert formula.expression == "=1+1"
    assert "formula cell(s) preserved without evaluation" in result.warnings[0]


def test_xlsx_external_reference_formula_is_not_followed(tool):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "formula"])
    sheet.append(["x", "='[external.xlsx]Sheet1'!A1"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    result = tool.read(
        raw_bytes=buffer.getvalue(),
        source_uri="file:///external.xlsx",
        file_format="xlsx",
    )
    formula = result.parsed_result.rows[0][1]
    assert isinstance(formula, FormulaCell)
    assert "[external.xlsx]" in formula.expression


def test_xlsx_missing_sheet_is_typed(tool):
    result = tool.read(
        raw_bytes=make_xlsx(),
        source_uri="file:///input.xlsx",
        file_format="xlsx",
        sheet_name="Missing",
    )
    assert result.error.category is ErrorCategory.SOURCE_NOT_FOUND


@pytest.mark.parametrize("raw", [b"not-a-zip", b"", b"PK\x03\x04broken"])
def test_corrupted_xlsx_is_rejected(tool, raw):
    result = tool.read(
        raw_bytes=raw,
        source_uri="file:///input.xlsx",
        file_format="xlsx",
    )
    assert result.error.category is ErrorCategory.MALFORMED_SOURCE


@pytest.mark.parametrize("file_format", ["xls", "xlsm", "ods", "json"])
def test_unsupported_formats_are_explicit(tool, file_format):
    result = tool.read(
        raw_bytes=b"content",
        source_uri=f"file:///input.{file_format}",
        file_format=file_format,
    )
    assert result.error.category is ErrorCategory.UNSUPPORTED_FORMAT


def test_input_bytes_remain_unchanged(tool):
    raw = make_xlsx()
    before = bytes(raw)
    tool.read(raw_bytes=raw, source_uri="file:///input.xlsx", file_format="xlsx")
    assert raw == before


def test_same_parsed_csv_with_byte_difference_has_different_checksum(tool):
    lf = b"name,value\na,1\n"
    crlf = b"name,value\r\na,1\r\n"
    left = tool.read(raw_bytes=lf, source_uri="file:///lf.csv", file_format="csv")
    right = tool.read(raw_bytes=crlf, source_uri="file:///crlf.csv", file_format="csv")
    assert left.parsed_result == right.parsed_result
    assert left.source_metadata.checksum_sha256 != right.source_metadata.checksum_sha256


def test_task_run_request_and_source_metadata(tool):
    task_id = generate_task_id()
    run_id = generate_run_id()
    result = tool.read(
        raw_bytes=b"name,value\na,1\n",
        source_uri="file:///input.csv",
        file_format="csv",
        task_id=task_id,
        run_id=run_id,
    )
    assert result.task_id == task_id
    assert result.run_id == run_id
    assert result.request_id == REQUEST_ID
    assert result.source_metadata.fetched_at == NOW
    assert result.source_metadata.data_as_of is None


def test_invalid_correlation_id_is_rejected(tool):
    with pytest.raises(ValueError):
        tool.read(
            raw_bytes=b"name,value\na,1\n",
            source_uri="file:///input.csv",
            file_format="csv",
            task_id="bad",
        )


def test_source_uri_credentials_are_rejected(tool):
    result = tool.read(
        raw_bytes=b"name,value\na,1\n",
        source_uri="https://user:password@example.com/input.csv",
        file_format="csv",
    )
    assert result.error.category is ErrorCategory.MALFORMED_INPUT


def test_result_repr_does_not_log_raw_csv(tool):
    result = tool.read(
        raw_bytes=b"name,value\nprivate-customer,1\n",
        source_uri="file:///input.csv",
        file_format="csv",
    )
    assert "private-customer" not in repr(result)
