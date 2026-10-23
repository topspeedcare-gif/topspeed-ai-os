"""Non-exported, caller-supplied-byte Coupang Wing Source-envelope foundation."""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from xml.etree.ElementTree import fromstring
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from topspeed_ai_os.core.canonical import payload_sha256

from ..base import InvocationContext, ReadTool
from ..checksum import raw_sha256
from ..errors import ToolFailure
from ..result import ErrorCategory as _Category
from ..result import ExecutionMode, SensitiveBytes, SourceMetadata, ToolResult, ToolStatus
from .tabular import TabularData

_SOURCE = "coupang_wing_export"
_TOOL_NAME = "coupang_wing_export_read"
_INVALID_OPERATION = "coupang_wing_export.read_invalid_request"
_SCHEMA_VERSION = "topspeed.coupang_wing_export.source_envelope.v1"
_SALES_SCHEMA_VERSION = "topspeed.coupang_wing_export.source_envelope.v2"
_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_WORKBOOK_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)
_SELECTION_REFERENCE = (
    "docs/architecture/"
    "ADR_PHASE5_COUPANG_PERIOD_AGGREGATE_READ_ONLY_INVENTORY_ASSESSMENT_V1.md"
    "@f5952d3dcb808c4cb65e6f18bc64b52afcbeab6d"
)
_SALES_SELECTION_REFERENCE = (
    "docs/architecture/"
    "ADR_PHASE5_COUPANG_L_REPEATABLE_COMPLETED_30D_READ_ONLY_AUTHORITY_V1.md"
    "@b02fb74a766f4f85e7c7d878a3a5faa4a16bc7ba"
)
_KST = timezone(timedelta(hours=9))
# Inventory V1 retains its published earliest-observation boundary. The
# assessment additionally enforces the caller's dynamic sales-period close.
_INVENTORY_V1_CLOSE = datetime(2026, 10, 20, tzinfo=_KST)
_CREATED_TAG = "{http://purl.org/dc/terms/}created"
_CONTENT_OVERRIDE = "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
_CREATED_PATTERN = re.compile(
    r"[0-9]{4}-[0-9]{2}-[0-9]{2}T[0-9]{2}:[0-9]{2}:[0-9]{2}"
    r"(?:Z|[+-](?:[01][0-9]|2[0-3]):[0-5][0-9])"
)


@dataclass(frozen=True, slots=True)
class _ExportSpec:
    source_uri: str
    worksheet: str
    operation: str
    headers: tuple[str, ...]
    first_data_row: int


_SALES = _ExportSpec(
    "coupang-wing://seller-insights/vendor-item-metrics",
    "vendor item metrics",
    "coupang_wing_export.read_sales_period_aggregate",
    ("등록상품ID", "옵션 ID", "판매방식", "총 판매수", "총 취소된 상품수", "판매량"),
    2,
)
_INVENTORY = _ExportSpec(
    "coupang-wing://inventory-health/sku-info",
    "Excel Data",
    "coupang_wing_export.read_inventory_health_observation",
    (
        "등록상품 ID",
        "옵션 ID",
        "SKU ID",
        "상품등급",
        "판매가능재고 (실시간 기준)",
        "입고예정재고(실시간 기준)",
        "고객반품 지난 30일(전일자 기준)",
    ),
    3,
)


def _select_spec(export_kind: str) -> _ExportSpec:
    if type(export_kind) is str:
        if export_kind == "SALES_PERIOD_AGGREGATE":
            return _SALES
        if export_kind == "INVENTORY_HEALTH_OBSERVATION":
            return _INVENTORY
    raise ToolFailure(_Category.MALFORMED_INPUT, "Export kind is invalid")


def _completed_period_close(period_start: str, period_end: str, period_timezone: str) -> datetime:
    """Validate explicit selection only; no workbook or acquisition-time inference."""
    if any(type(value) is not str for value in (period_start, period_end, period_timezone)):
        raise TypeError("Period fields must be exact strings")
    if period_timezone != "Asia/Seoul":
        raise ValueError("Period timezone is not authorized")
    for value in (period_start, period_end):
        if re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value) is None:
            raise ValueError("Period date is not canonical")
    start, end = date.fromisoformat(period_start), date.fromisoformat(period_end)
    if start.isoformat() != period_start or end.isoformat() != period_end:
        raise ValueError("Period date is not canonical")
    if (end - start).days != 29:
        raise ValueError("Period must contain exactly thirty inclusive days")
    try:
        close_date = end + timedelta(days=1)
    except OverflowError as exc:
        raise ValueError("Period close is outside the supported calendar") from exc
    return datetime(close_date.year, close_date.month, close_date.day, tzinfo=_KST)


def _validate_request(
    raw_bytes: bytes,
    spec: _ExportSpec,
    period_start: str | None,
    period_end: str | None,
    period_timezone: str | None,
) -> datetime | None:
    if type(raw_bytes) is not bytes:
        raise ToolFailure(_Category.MALFORMED_INPUT, "Original input must be exact bytes")
    selection = (period_start, period_end, period_timezone)
    if spec is _SALES:
        try:
            return _completed_period_close(*selection)
        except (TypeError, ValueError, OverflowError) as exc:
            raise ToolFailure(
                _Category.MALFORMED_INPUT, "Sales selection is not authorized"
            ) from exc
    elif any(value is not None for value in selection):
        raise ToolFailure(_Category.MALFORMED_INPUT, "Inventory does not accept period fields")
    return None


def _workbook_created_at(raw_bytes: bytes) -> datetime:
    try:
        with ZipFile(io.BytesIO(raw_bytes)) as archive:
            if archive.namelist().count("docProps/core.xml") != 1:
                raise ToolFailure(
                    _Category.MALFORMED_SOURCE, "Workbook creation metadata is missing"
                )
            content_types = fromstring(archive.read("[Content_Types].xml"))
            workbook_types = [
                node.get("ContentType")
                for node in content_types.iter(_CONTENT_OVERRIDE)
                if node.get("PartName") == "/xl/workbook.xml"
            ]
            if len(workbook_types) != 1:
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook package is malformed")
            if workbook_types[0] != _WORKBOOK_CONTENT_TYPE:
                raise ToolFailure(_Category.UNSUPPORTED_FORMAT, "Workbook format is unsupported")
            core = fromstring(archive.read("docProps/core.xml"))
            created_nodes = list(core.iter(_CREATED_TAG))
            if len(created_nodes) != 1 or len(created_nodes[0]):
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook creation time is ambiguous")
            value = created_nodes[0].text
            if value is None or _CREATED_PATTERN.fullmatch(value) is None:
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook creation time is invalid")
            return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(_KST)
    except (BadZipFile, KeyError, OSError, ValueError, SyntaxError, EOFError, OverflowError) as exc:
        raise ToolFailure(
            _Category.MALFORMED_SOURCE, "Workbook package or time is malformed"
        ) from exc


def _parse_worksheet(raw_bytes: bytes, spec: _ExportSpec) -> TabularData:
    try:
        workbook = load_workbook(
            io.BytesIO(raw_bytes), read_only=True, data_only=False, keep_links=False
        )
        try:
            if spec.worksheet not in workbook.sheetnames:
                raise ToolFailure(_Category.SOURCE_NOT_FOUND, "Required worksheet is absent")
            worksheet = workbook[spec.worksheet]
            # Declared dimensions must not pad, truncate, or conceal actual row widths.
            worksheet.reset_dimensions()
            rows = worksheet.iter_rows()
            header = next(rows, ())
            if any(cell.data_type == "f" for cell in header):
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Formula header is not permitted")
            headers = tuple(str(cell.value) for cell in header)
            if any(headers.count(required) != 1 for required in spec.headers):
                raise ToolFailure(_Category.SOURCE_SCHEMA_CHANGED, "Required headers have changed")
            if spec.first_data_row == 3:
                next(rows, None)  # The explanatory second row is not a data record.
            body = []
            for row in rows:
                if len(row) != len(headers):
                    raise ToolFailure(
                        _Category.SOURCE_SCHEMA_CHANGED, "Worksheet row shape changed"
                    )
                # All returned body cells are consumed by this byte-to-tabular boundary.
                if any(cell.data_type == "f" for cell in row):
                    raise ToolFailure(_Category.MALFORMED_SOURCE, "Formula data is not permitted")
                body.append(tuple(cell.value for cell in row))
            if not body or all(all(value is None for value in row) for row in body):
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Worksheet data body is empty")
            return TabularData(headers, tuple(body), spec.worksheet)
        finally:
            workbook.close()
    except InvalidFileException as exc:
        raise ToolFailure(_Category.UNSUPPORTED_FORMAT, "Workbook format is unsupported") from exc
    except (IndexError, TypeError) as exc:
        # Malformed string/style references and typed XML fields can fail inside
        # openpyxl decoding. An injected or local implementation fault is not
        # source evidence merely because it has the same exception type.
        decoder_trace = exc.__traceback__
        while decoder_trace is not None and decoder_trace.tb_next is not None:
            decoder_trace = decoder_trace.tb_next
        origin = (
            decoder_trace.tb_frame.f_globals.get("__name__", "")
            if decoder_trace is not None
            else ""
        )
        if not origin.startswith(
            (
                "openpyxl.reader.",
                "openpyxl.worksheet.",
                "openpyxl.descriptors.",
                "openpyxl.styles.",
            )
        ):
            raise
        raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook content is malformed") from exc
    except (BadZipFile, KeyError, OSError, ValueError, SyntaxError, EOFError) as exc:
        raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook content is malformed") from exc


def _source_metadata(
    *,
    spec: _ExportSpec,
    export_kind: str,
    checksum: str,
    created_at: datetime,
    fetched_at: datetime,
    period_start: str | None,
    period_end: str | None,
    period_timezone: str | None,
    close: datetime | None,
) -> SourceMetadata:
    sales = spec is _SALES
    data_as_of = close if sales else created_at
    period_close = close.isoformat() if sales else None
    schema = _SALES_SCHEMA_VERSION if sales else _SCHEMA_VERSION
    authority = _SALES_SELECTION_REFERENCE if sales else _SELECTION_REFERENCE
    created_text = created_at.isoformat()
    envelope = {
        "schema_version": schema,
        "source": _SOURCE,
        "source_uri": spec.source_uri,
        "export_kind": export_kind,
        "worksheet_name": spec.worksheet,
        "raw_checksum_sha256": checksum,
        "workbook_created_at": created_text,
        "period_start": period_start,
        "period_end": period_end,
        "period_timezone": period_timezone,
        "period_close": period_close,
        "data_as_of": data_as_of.isoformat(),
        "selection_authority_reference": authority,
    }
    # Fresh scalar-only dictionary: no producer reference or shared mutable backing survives.
    return SourceMetadata(
        source=_SOURCE,
        source_uri=spec.source_uri,
        fetched_at=fetched_at,
        data_as_of=data_as_of,
        checksum_sha256=checksum,
        content_type=_CONTENT_TYPE,
        http_status=None,
        modified_time=None,
        source_specific_metadata={
            "schema_version": schema,
            "export_kind": export_kind,
            "worksheet_name": spec.worksheet,
            "workbook_created_at": created_text,
            "raw_artifact_checksum_sha256": checksum,
            "original_source_uri": spec.source_uri,
            "period_start": period_start,
            "period_end": period_end,
            "period_timezone": period_timezone,
            "period_close": period_close,
            "selection_authority_reference": authority,
            "envelope_sha256": payload_sha256(envelope),
        },
    )


class CoupangWingExportReadTool(ReadTool):
    def __init__(
        self,
        *,
        clock: Callable[[], datetime],
        request_id_factory: Callable[[], str] | None = None,
    ) -> None:
        if not callable(clock):
            raise TypeError("An explicit callable clock is required")
        super().__init__(
            tool_name=_TOOL_NAME,
            _execution_mode=ExecutionMode.MOCK,
            clock=clock,
            request_id_factory=request_id_factory,
        )
        self._clock = clock  # Also preserve an explicitly supplied falsey callable.

    def read(
        self,
        *,
        raw_bytes: bytes,
        export_kind: str,
        period_start: str | None = None,
        period_end: str | None = None,
        period_timezone: str | None = None,
        task_id: str | None = None,
        run_id: str | None = None,
    ) -> ToolResult:
        # Existing correlation-ID, request-ID factory, and clock exceptions propagate.
        context = self._prepare_invocation(task_id=task_id, run_id=run_id)
        fetched_at = self._now()
        operation = _INVALID_OPERATION
        try:
            spec = _select_spec(export_kind)
            operation = spec.operation
            close = _validate_request(raw_bytes, spec, period_start, period_end, period_timezone)
            checksum = raw_sha256(raw_bytes)
            created_at = _workbook_created_at(raw_bytes)
            earliest = close if spec is _SALES else _INVENTORY_V1_CLOSE
            if created_at < earliest or fetched_at < created_at:
                raise ToolFailure(_Category.MALFORMED_SOURCE, "Workbook source time is ineligible")
            parsed = _parse_worksheet(raw_bytes, spec)
            metadata = _source_metadata(
                spec=spec,
                export_kind=export_kind,
                checksum=checksum,
                created_at=created_at,
                fetched_at=fetched_at,
                period_start=period_start,
                period_end=period_end,
                period_timezone=period_timezone,
                close=close,
            )
        except ToolFailure as failure:
            return self._failed(operation, context, failure)
        except Exception:
            return self._failed(
                operation,
                context,
                ToolFailure(_Category.UNKNOWN, "Source envelope validation failed"),
            )
        return ToolResult(
            tool_name=_TOOL_NAME,
            operation=operation,
            request_id=context.request_id,
            status=ToolStatus.SUCCESS,
            execution_mode=ExecutionMode.MOCK,
            task_id=context.task_id,
            run_id=context.run_id,
            source_metadata=metadata,
            raw_result=SensitiveBytes(raw_bytes),
            parsed_result=parsed,
            error=None,
            attempt_count=1,
            warnings=(),
        )

    @staticmethod
    def _failed(operation: str, context: InvocationContext, failure: ToolFailure) -> ToolResult:
        return ToolResult(
            tool_name=_TOOL_NAME,
            operation=operation,
            request_id=context.request_id,
            status=ToolStatus.FAILED,
            execution_mode=ExecutionMode.MOCK,
            task_id=context.task_id,
            run_id=context.run_id,
            source_metadata=None,
            raw_result=None,
            parsed_result=None,
            error=failure.to_error(1),
            attempt_count=1,
            warnings=(),
        )
