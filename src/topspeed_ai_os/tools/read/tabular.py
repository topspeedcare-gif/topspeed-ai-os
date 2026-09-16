"""Pure read-only CSV and XLSX parsers over caller-supplied raw bytes."""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ..base import ReadTool
from ..checksum import raw_sha256
from ..errors import ToolFailure
from ..redaction import sanitized_source_uri
from ..result import (
    ErrorCategory,
    ExecutionMode,
    SensitiveBytes,
    SourceMetadata,
    ToolResult,
    ToolStatus,
)


@dataclass(frozen=True, slots=True)
class FormulaCell:
    expression: str


@dataclass(frozen=True, slots=True)
class TabularData:
    headers: tuple[str, ...]
    rows: tuple[tuple[Any, ...], ...]
    sheet_name: str | None


class TabularReadTool(ReadTool):
    def __init__(
        self,
        *,
        execution_mode: ExecutionMode,
        clock: Callable[[], datetime] | None = None,
        request_id_factory: Callable[[], str] | None = None,
    ) -> None:
        super().__init__(
            tool_name="tabular_read",
            execution_mode=execution_mode,
            clock=clock,
            request_id_factory=request_id_factory,
        )

    def read(
        self,
        *,
        raw_bytes: bytes,
        source_uri: str,
        file_format: str,
        encoding: str = "utf-8",
        delimiter: str = ",",
        sheet_name: str | None = None,
        task_id: str | None = None,
        run_id: str | None = None,
    ) -> ToolResult:
        context = self._prepare_invocation(task_id=task_id, run_id=run_id)
        operation = f"tabular.read_{file_format.lower()}"
        fetched_at = self._now()
        checksum = raw_sha256(raw_bytes)
        try:
            safe_uri = sanitized_source_uri(source_uri)
        except ValueError:
            return self._failed(
                operation,
                context,
                ToolFailure(ErrorCategory.MALFORMED_INPUT, "Source URI is unsafe"),
            )
        normalized_format = file_format.lower().lstrip(".")
        try:
            if normalized_format == "csv":
                parsed, warnings = self._parse_csv(raw_bytes, encoding, delimiter)
            elif normalized_format == "xlsx":
                parsed, warnings = self._parse_xlsx(raw_bytes, sheet_name)
            else:
                raise ToolFailure(
                    ErrorCategory.UNSUPPORTED_FORMAT,
                    "Tabular format is not supported",
                )
        except ToolFailure as failure:
            return self._failed(operation, context, failure)
        metadata = SourceMetadata(
            source="tabular_file",
            source_uri=safe_uri,
            fetched_at=fetched_at,
            data_as_of=None,
            checksum_sha256=checksum,
            content_type=(
                "text/csv"
                if normalized_format == "csv"
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            source_specific_metadata={"file_name": PurePosixPath(safe_uri).name},
        )
        return ToolResult(
            tool_name=self.tool_name,
            operation=operation,
            request_id=context.request_id,
            status=ToolStatus.SUCCESS,
            execution_mode=self.execution_mode,
            task_id=context.task_id,
            run_id=context.run_id,
            source_metadata=metadata,
            raw_result=SensitiveBytes(raw_bytes),
            parsed_result=parsed,
            attempt_count=1,
            warnings=warnings,
        )

    @staticmethod
    def _parse_csv(
        raw_bytes: bytes,
        encoding: str,
        delimiter: str,
    ) -> tuple[TabularData, tuple[str, ...]]:
        if not raw_bytes:
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "CSV source is empty")
        if len(delimiter) != 1:
            raise ToolFailure(ErrorCategory.MALFORMED_INPUT, "CSV delimiter must be one character")
        try:
            text = raw_bytes.decode(encoding, errors="strict")
        except (LookupError, UnicodeDecodeError) as exc:
            raise ToolFailure(
                ErrorCategory.MALFORMED_SOURCE,
                "CSV encoding is invalid",
            ) from exc
        try:
            rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True))
        except csv.Error as exc:
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "CSV source is malformed") from exc
        if not rows or not rows[0] or all(not value.strip() for value in rows[0]):
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "CSV header is missing")
        headers = tuple(rows[0])
        if len(set(headers)) != len(headers):
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "CSV header contains duplicates")
        for row in rows[1:]:
            if len(row) != len(headers):
                raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "CSV row width is malformed")
        return TabularData(headers, tuple(tuple(row) for row in rows[1:]), None), ()

    @staticmethod
    def _parse_xlsx(
        raw_bytes: bytes,
        sheet_name: str | None,
    ) -> tuple[TabularData, tuple[str, ...]]:
        if not raw_bytes:
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "XLSX source is empty")
        try:
            workbook = load_workbook(
                io.BytesIO(raw_bytes),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
            raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "XLSX source is malformed") from exc
        try:
            selected = sheet_name or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise ToolFailure(ErrorCategory.SOURCE_NOT_FOUND, "XLSX sheet was not found")
            worksheet = workbook[selected]
            values: list[tuple[Any, ...]] = []
            formula_count = 0
            for row in worksheet.iter_rows():
                converted: list[Any] = []
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
                        converted.append(FormulaCell(str(cell.value)))
                    else:
                        converted.append(cell.value)
                values.append(tuple(converted))
            if not values or not values[0] or all(value in {None, ""} for value in values[0]):
                raise ToolFailure(ErrorCategory.MALFORMED_SOURCE, "XLSX header is missing")
            headers = tuple(str(value) for value in values[0])
            if len(set(headers)) != len(headers):
                raise ToolFailure(
                    ErrorCategory.MALFORMED_SOURCE,
                    "XLSX header contains duplicates",
                )
            warnings = (
                (f"{formula_count} formula cell(s) preserved without evaluation",)
                if formula_count
                else ()
            )
            return TabularData(headers, tuple(values[1:]), selected), warnings
        finally:
            workbook.close()

    def _failed(
        self,
        operation: str,
        context: Any,
        failure: ToolFailure,
    ) -> ToolResult:
        return ToolResult(
            tool_name=self.tool_name,
            operation=operation,
            request_id=context.request_id,
            status=ToolStatus.FAILED,
            execution_mode=self.execution_mode,
            task_id=context.task_id,
            run_id=context.run_id,
            source_metadata=None,
            error=failure.to_error(1),
            attempt_count=1,
        )
